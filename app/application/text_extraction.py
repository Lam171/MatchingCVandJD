"""Text extraction adapters used by the upload API.

Keep this module independent from FastAPI so each parser can be tested directly.
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path


class TextExtractionError(Exception):
    """A user-safe parsing error with an API error code."""

    def __init__(self, code: str, message: str, status_code: int = 422) -> None:
        self.code = code
        self.message = message
        self.status_code = status_code
        super().__init__(message)


SUPPORTED_SUFFIXES = {".txt", ".pdf", ".docx"}
MIN_TEXT_LENGTH = 20


def extract_text(payload: bytes, filename: str) -> tuple[str, str]:
    """Extract text and report the method used for a supported uploaded file."""
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise TextExtractionError(
            "UNSUPPORTED_FILE_TYPE",
            "Chỉ hỗ trợ TXT, PDF có text hoặc Word (.docx).",
            415,
        )
    if not payload:
        raise TextExtractionError("EMPTY_FILE", "Tệp tải lên đang trống.")

    if suffix == ".txt":
        return _extract_txt(payload), "plain_text"
    if suffix == ".docx":
        return _extract_docx(payload), "docx"
    if suffix == ".pdf":
        return _extract_pdf(payload)


def _extract_txt(payload: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "cp1258"):
        try:
            text = payload.decode(encoding).strip()
            if text:
                return text
        except UnicodeDecodeError:
            continue
    raise TextExtractionError("INVALID_TEXT_ENCODING", "Không thể đọc mã hoá của tệp TXT.")


def _extract_docx(payload: bytes) -> str:
    try:
        from docx import Document
    except ImportError as error:
        raise TextExtractionError(
            "PARSER_NOT_INSTALLED", "Thiếu thư viện python-docx trên máy chủ.", 503
        ) from error

    try:
        document = Document(BytesIO(payload))
        parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            parts.extend(cell.text.strip() for row in table.rows for cell in row.cells if cell.text.strip())
    except Exception as error:
        raise TextExtractionError("INVALID_DOCX", "Không thể đọc tệp Word (.docx).") from error
    return _require_text("\n".join(parts), "Tệp Word không chứa nội dung văn bản.")


def _extract_pdf(payload: bytes) -> tuple[str, str]:
    try:
        import fitz
    except ImportError as error:
        raise TextExtractionError("PARSER_NOT_INSTALLED", "Thiếu thư viện PyMuPDF trên máy chủ.", 503) from error

    try:
        document = fitz.open(stream=payload, filetype="pdf")
        text = "\n".join(page.get_text("text") for page in document).strip()
        if len(text) >= MIN_TEXT_LENGTH:
            return text, "pdf_text"
    except TextExtractionError:
        raise
    except Exception as error:
        raise TextExtractionError("INVALID_PDF", "Không thể đọc tệp PDF.") from error

    raise TextExtractionError(
        "PDF_SCAN_NOT_SUPPORTED",
        "PDF không có lớp text (có thể là PDF scan). Phiên bản này không hỗ trợ OCR.",
    )


def _require_text(text: str, message: str) -> str:
    cleaned = text.strip()
    if len(cleaned) < MIN_TEXT_LENGTH:
        raise TextExtractionError("TEXT_TOO_SHORT", message)
    return cleaned
